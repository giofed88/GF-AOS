import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document


SCRIPT = Path(__file__).parents[1] / "scripts" / "build_document_style_template.py"
SPEC = importlib.util.spec_from_file_location("document_style", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class DocumentStyleTest(unittest.TestCase):
    def test_builds_all_profiles_as_a4_with_stable_identity(self):
        with tempfile.TemporaryDirectory() as temporary:
            for name, profile in MODULE.PROFILES.items():
                output = Path(temporary) / f"{name}.docx"
                MODULE.build(output, name)
                document = Document(output)
                section = document.sections[0]
                self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
                self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn(profile.title, text)
                self.assertIn("BOZZA DA VALIDARE", text)
                with zipfile.ZipFile(output) as archive:
                    header = archive.read("word/header1.xml").decode("utf-8")
                    relations = archive.read("word/_rels/header1.xml.rels").decode("utf-8")
                self.assertIn(profile.label, header)
                self.assertIn("image", relations)

    def test_non_engagement_profiles_are_sanitized_drafts(self):
        with tempfile.TemporaryDirectory() as temporary:
            for name in ("professional-report", "minutes", "working-paper"):
                output = Path(temporary) / f"{name}.docx"
                MODULE.build(output, name)
                document = Document(output)
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn("BOZZA DA VALIDARE", text)
                self.assertIn("[DENOMINAZIONE CLIENTE]", "\n".join(
                    cell.text for table in document.tables for row in table.rows for cell in row.cells
                ))

    def test_unknown_profile_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaises(KeyError):
                MODULE.build(Path(temporary) / "unknown.docx", "unknown")


if __name__ == "__main__":
    unittest.main()
