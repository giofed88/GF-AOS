import importlib.util
import json
import tempfile
import unittest
from pathlib import Path


SCRIPT = Path(__file__).parents[1] / "scripts" / "document_intelligence.py"
SPEC = importlib.util.spec_from_file_location("document_intelligence", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class DocumentIntelligenceTest(unittest.TestCase):
    def test_extracts_supported_documents_and_preserves_sources(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source"
            output = root / "workspace"
            source.mkdir()

            (source / "nota.txt").write_text("Evidenza testuale", encoding="utf-8")
            (source / "dati.csv").write_text("voce;valore\nricavi;100\n", encoding="utf-8")

            doc = Document()
            doc.add_paragraph("Verbale di prova")
            doc.save(source / "verbale.docx")

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Situazione"
            sheet["A1"] = "Totale"
            sheet["B1"] = "=SUM(B2:B3)"
            workbook.save(source / "situazione.xlsx")

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Conclusioni"
            presentation.save(source / "presentazione.pptx")

            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            with (source / "allegato.pdf").open("wb") as handle:
                writer.write(handle)

            before = {path.name: MODULE.sha256_file(path) for path in source.iterdir()}
            result = MODULE.build_index(source, output)
            after = {path.name: MODULE.sha256_file(path) for path in source.iterdir()}

            self.assertEqual(before, after)
            self.assertEqual(result["count"], 6)
            self.assertEqual(result["errors"], 0)
            self.assertTrue((output / "document_index.json").is_file())
            self.assertTrue((output / "document_index.md").is_file())
            saved = json.loads((output / "document_index.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["source_mode"], "read_only")
            spreadsheet = next(item for item in saved["records"] if item["path"] == "situazione.xlsx")
            self.assertEqual(spreadsheet["details"]["formula_cells"], 1)
            self.assertEqual(spreadsheet["text_format"], "markdown")
            self.assertTrue(spreadsheet["text_locator"].endswith(".md"))

    def test_rejects_output_inside_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source"
            source.mkdir()
            with self.assertRaisesRegex(ValueError, "esterno"):
                MODULE.build_index(source, source / "output")

    def test_rejects_symbolic_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            target = root / "target"
            target.mkdir()
            link = root / "link"
            link.symlink_to(target, target_is_directory=True)
            with self.assertRaisesRegex(ValueError, "simbolica"):
                MODULE.build_index(link, root / "output")


if __name__ == "__main__":
    unittest.main()
