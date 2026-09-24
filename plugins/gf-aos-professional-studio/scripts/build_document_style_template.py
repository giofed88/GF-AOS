#!/usr/bin/env python3
"""Genera il riferimento DOCX sanificato del design system GF-AOS."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import NamedTuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


INK = "20252B"
PETROL = "245865"
PETROL_LIGHT = "E8F0F1"
WARM = "F4F1EB"
LINE = "D9D6CF"
WHITE = "FFFFFF"
HEADER_LEFT = Path(__file__).resolve().parents[1] / "assets" / "document" / "gf-header-left-v1.png"


class DocumentProfile(NamedTuple):
    label: str
    title: str
    subtitle: str
    sections: tuple[tuple[str, str], ...]


PROFILES = {
    "engagement-letter": DocumentProfile(
        "LETTERA DI INCARICO",
        "Lettera di incarico professionale",
        "[AREA PROFESSIONALE]  ·  [PERIODO DI RIFERIMENTO]",
        (),
    ),
    "professional-report": DocumentProfile(
        "RELAZIONE PROFESSIONALE",
        "Relazione professionale",
        "[OGGETTO]  ·  [PERIODO DI RIFERIMENTO]",
        (
            ("Premessa e mandato", "[CONTESTO, QUESITO E LIMITI DELL’INCARICO]"),
            ("Documentazione esaminata", "[FONTI INTERNE, EVIDENZE E DATA DI VERIFICA]"),
            ("Analisi", "[FATTI ACCERTATI, VALUTAZIONI E IPOTESI TENUTI DISTINTI]"),
            ("Conclusioni", "[ESITO, LIMITI, RISERVE E AZIONI RACCOMANDATE]"),
        ),
    ),
    "minutes": DocumentProfile(
        "VERBALE",
        "Verbale professionale",
        "[ORGANO / RIUNIONE]  ·  [DATA E ORA]",
        (
            ("Costituzione e presenze", "[LUOGO, MODALITÀ, PRESENTI, ASSENTI E DELEGHE]"),
            ("Ordine del giorno", "[PUNTI SOTTOPOSTI ALL’ESAME]"),
            ("Discussione", "[INTERVENTI E DOCUMENTI RICHIAMATI IN FORMA PERTINENTE]"),
            ("Deliberazioni e azioni", "[DECISIONI, RESPONSABILI, TERMINI E RISERVE]"),
        ),
    ),
    "working-paper": DocumentProfile(
        "CARTA DI LAVORO",
        "Carta di lavoro",
        "[AREA / CICLO]  ·  [PERIODO DI RIFERIMENTO]",
        (
            ("Obiettivo e asserzioni", "[OBIETTIVO DELLA PROCEDURA E ASSERZIONI RILEVANTI]"),
            ("Popolazione e campione", "[ORIGINE, PERIMETRO, CRITERIO E DIMENSIONE]"),
            ("Procedura ed evidenze", "[ATTIVITÀ SVOLTE, LOCATOR E RISULTATI VERIFICABILI]"),
            ("Esito e riesame", "[CONCLUSIONE, ECCEZIONI, FOLLOW-UP, PREPARATORE E RIESAMINATORE]"),
        ),
    ),
}


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color: str = LINE, size: str = "5") -> None:
    props = table._tbl.tblPr
    current = props.first_child_found_in("w:tblBorders")
    if current is None:
        current = OxmlElement("w:tblBorders")
        props.append(current)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = current.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            current.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def no_borders(table) -> None:
    borders(table, WHITE, "0")


def keep(paragraph, *, next_paragraph=False) -> None:
    props = paragraph._p.get_or_add_pPr()
    if props.find(qn("w:keepLines")) is None:
        props.append(OxmlElement("w:keepLines"))
    if next_paragraph and props.find(qn("w:keepNext")) is None:
        props.append(OxmlElement("w:keepNext"))


def set_repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:tblHeader"))


def page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGINA ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(PETROL)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def header(section, document_label: str) -> None:
    area = section.header
    area.is_linked_to_previous = False
    table = area.add_table(rows=1, cols=2, width=Cm(16.6))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(10.2)
    table.columns[1].width = Cm(6.4)
    no_borders(table)

    # The approved PDF composition is inserted as one high-resolution unit.
    # This prevents Word from independently redistributing monogram, name,
    # qualifications and decorative rule across table columns.
    identity = table.cell(0, 0)
    margins(identity, 0, 0, 0, 60)
    p = identity.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(HEADER_LEFT), width=Cm(8.45))

    right = table.cell(0, 1)
    margins(right, 50, 80, 50, 40)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{document_label}\n[DATA]  ·  REV [00]")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(PETROL)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def footer(section) -> None:
    area = section.footer
    area.is_linked_to_previous = False
    table = area.add_table(rows=1, cols=2, width=Cm(16.6))
    table.columns[0].width = Cm(12.8)
    table.columns[1].width = Cm(3.8)
    no_borders(table)
    p = table.cell(0, 0).paragraphs[0]
    r = p.add_run("RISERVATO  ·  USO PROFESSIONALE")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string("666666")
    page_number(table.cell(0, 1).paragraphs[0])


def configure(document: Document, document_label: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.7)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.7)
    header(section, document_label)
    footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    title = document.styles["Title"]
    title.font.name = "Georgia"
    title.font.size = Pt(23)
    title.font.bold = False
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(5)

    for name, size in (("Heading 1", 12.5), ("Heading 2", 10.5)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(PETROL)
        style.paragraph_format.space_before = Pt(11 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(4)


def add_section(document: Document, number: str, title: str, text: str) -> None:
    heading = document.add_heading(f"{number}  {title}", level=1)
    keep(heading, next_paragraph=True)
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.22)
    paragraph.add_run(text)


def add_profile_metadata(document: Document) -> None:
    meta = document.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta.columns[0].width = Cm(4.0)
    meta.columns[1].width = Cm(12.2)
    no_borders(meta)
    for index, (label, value) in enumerate((
        ("Cliente", "[DENOMINAZIONE CLIENTE]"),
        ("Incarico", "[DESCRIZIONE SINTETICA DELL’INCARICO]"),
        ("Riferimento", "[ID INCARICO]  ·  [DATA]"),
    )):
        meta.cell(index, 0).text = label.upper()
        meta.cell(index, 1).text = value
        margins(meta.cell(index, 0), 60, 0, 60, 120)
        margins(meta.cell(index, 1), 60, 0, 60, 0)
        for run in meta.cell(index, 0).paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(7.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(PETROL)
        for run in meta.cell(index, 1).paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)
    document.add_paragraph()


def add_profile_cover(document: Document, profile: DocumentProfile) -> None:
    title = document.add_paragraph(style="Title")
    title.add_run(profile.title)
    subtitle = document.add_paragraph(profile.subtitle)
    subtitle.paragraph_format.space_after = Pt(16)
    for run in subtitle.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(PETROL)
    add_profile_metadata(document)


def build_generic(document: Document, profile: DocumentProfile) -> None:
    add_profile_cover(document, profile)
    document.add_paragraph(
        "BOZZA DA VALIDARE · Il contenuto usa esclusivamente segnaposto e deve essere "
        "completato con fonti verificate e dati strettamente necessari."
    )
    for index, (title, text) in enumerate(profile.sections, start=1):
        add_section(document, str(index), title, text)
    add_section(
        document,
        str(len(profile.sections) + 1),
        "Riservatezza e approvazione",
        "[DATI MINIMIZZATI, ACCESSI AUTORIZZATI, TEMPI DI CONSERVAZIONE E GATE DI APPROVAZIONE]",
    )


def build(output: Path, profile_name: str = "engagement-letter") -> None:
    profile = PROFILES[profile_name]
    document = Document()
    configure(document, profile.label)

    if profile_name != "engagement-letter":
        build_generic(document, profile)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        return

    add_profile_cover(document, profile)

    draft = document.add_paragraph("BOZZA DA VALIDARE")
    draft.paragraph_format.space_after = Pt(10)
    for run in draft.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(PETROL)

    intro = document.add_paragraph(
        "Gentile [NOME CLIENTE], con la presente definiamo in modo trasparente il perimetro "
        "dell’incarico, le responsabilità reciproche e le condizioni economiche applicabili."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.space_after = Pt(12)

    add_section(document, "1", "Oggetto e finalità", "L’incarico riguarda [DESCRIZIONE]. Il risultato atteso è [OUTPUT], destinato a [DESTINATARIO E FINALITÀ].")
    add_bullet(document, "attività comprese: [ELENCO SINTETICO E VERIFICABILE]")
    add_bullet(document, "attività escluse: [ESCLUSIONI O PRESTAZIONI AGGIUNTIVE]")
    add_bullet(document, "periodo e scadenze: [PERIODO] · [SCADENZE CONFERMATE]")

    add_section(document, "2", "Metodo di lavoro e responsabilità", "Lo Studio opera con diligenza professionale, mantiene separate le valutazioni dalle evidenze e segnala tempestivamente limiti, documenti mancanti e circostanze che richiedono un approfondimento.")
    document.add_heading("Responsabilità del cliente", level=2)
    add_bullet(document, "fornire documenti completi, autentici e tempestivi;")
    add_bullet(document, "segnalare variazioni rilevanti per iscritto;")
    add_bullet(document, "approvare espressamente decisioni, invii e depositi quando richiesto.")

    add_section(document, "3", "Documenti e limiti", "La documentazione richiesta è elencata nell’allegato operativo. L’assenza o la tardiva consegna di elementi essenziali può limitare o impedire lo svolgimento dell’incarico e viene riportata nelle conclusioni.")

    # Keep the economic block together: a deliberate page transition is more
    # professional than a fee table split across two pages.
    document.add_page_break()
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    add_section(document, "4", "Compenso e condizioni", "Il corrispettivo è determinato in relazione a perimetro, complessità, responsabilità e tempi stimati. Le prestazioni ulteriori richiedono una proposta separata.")
    fees = document.add_table(rows=1, cols=4)
    fees.alignment = WD_TABLE_ALIGNMENT.CENTER
    fees.autofit = False
    widths = (Cm(7.2), Cm(3.1), Cm(2.8), Cm(3.0))
    for index, width in enumerate(widths):
        fees.columns[index].width = width
    headers = ("Prestazione", "Periodicità", "Corrispettivo", "Note")
    for index, label in enumerate(headers):
        cell = fees.cell(0, index)
        cell.text = label
        shade(cell, PETROL)
        margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_header(fees.rows[0])
    for row_index, values in enumerate((
        ("[PRESTAZIONE PRINCIPALE]", "[PERIODO]", "€ [IMPORTO]", "[ONERI / IVA]"),
        ("[PRESTAZIONE ACCESSORIA]", "[EVENTO]", "€ [IMPORTO]", "[CONDIZIONI]"),
        ("[ATTIVITÀ STRAORDINARIE]", "Su richiesta", "Da concordare", "Preventivo dedicato"),
    ), start=1):
        cells = fees.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            margins(cells[index])
            if row_index % 2 == 0:
                shade(cells[index], WARM)
            for run in cells[index].paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.8)
        for index in (1, 2):
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    borders(fees)

    add_section(document, "5", "Durata recesso e sospensione", "L’incarico decorre da [DATA] e termina il [DATA / EVENTO]. Recesso, sospensione per mancata collaborazione e attività di chiusura sono disciplinati dalle condizioni concordate e dalla normativa applicabile.")
    add_section(document, "6", "Riservatezza e dati personali", "I dati sono trattati per finalità connesse all’incarico e agli obblighi di legge, secondo basi giuridiche, destinatari e tempi di conservazione indicati nell’informativa allegata. Il consenso è richiesto soltanto quando costituisce la base giuridica appropriata.")
    add_bullet(document, "minimizzazione dei dati e accessi limitati alle persone autorizzate;")
    add_bullet(document, "conservazione per il tempo necessario o previsto dalla legge;")
    add_bullet(document, "diritti dell’interessato e reclamo all’autorità di controllo.")

    add_section(document, "7", "Approvazioni e comunicazioni", "Bozze, invii esterni, depositi e modifiche delle fonti restano distinti. Nessun documento è considerato approvato o trasmesso senza il relativo passaggio di conferma.")
    add_section(document, "8", "Accettazione", "Le parti dichiarano di aver letto e compreso il perimetro dell’incarico, le condizioni economiche, le responsabilità e gli allegati richiamati.")

    signatures = document.add_table(rows=2, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.autofit = False
    signatures.columns[0].width = Cm(8.0)
    signatures.columns[1].width = Cm(8.0)
    no_borders(signatures)
    labels = (("Il Cliente", "Il Professionista"), ("Data e firma ____________________", "Data e firma ____________________"))
    for row_index, row in enumerate(labels):
        for column, value in enumerate(row):
            cell = signatures.cell(row_index, column)
            cell.text = value
            margins(cell, 130 if row_index == 0 else 210, 100, 70, 100)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.8 if row_index else 9.5)
                run.font.bold = row_index == 0
                run.font.color.rgb = RGBColor.from_string(PETROL if row_index == 0 else INK)

    document.add_page_break()
    heading = document.add_paragraph(style="Title")
    heading.add_run("Allegato operativo")
    document.add_paragraph("Documenti richiesti  ·  responsabilità  ·  scadenze")
    appendix = document.add_table(rows=1, cols=4)
    appendix.alignment = WD_TABLE_ALIGNMENT.CENTER
    appendix.autofit = False
    for index, width in enumerate((Cm(6.2), Cm(3.2), Cm(3.2), Cm(3.5))):
        appendix.columns[index].width = width
    for index, label in enumerate(("Documento o attività", "Responsabile", "Termine", "Stato")):
        cell = appendix.cell(0, index)
        cell.text = label
        shade(cell, PETROL)
        margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_header(appendix.rows[0])
    for row_index in range(1, 7):
        cells = appendix.add_row().cells
        values = (f"[ELEMENTO {row_index}]", "[RUOLO]", "[DATA]", "Da acquisire")
        for column, value in enumerate(values):
            cells[column].text = value
            margins(cells[column], 130, 100, 130, 100)
            if row_index % 2 == 0:
                shade(cells[column], PETROL_LIGHT)
            for run in cells[column].paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.8)
    borders(appendix)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build GF-AOS document style reference")
    parser.add_argument("output")
    parser.add_argument("--profile", choices=sorted(PROFILES), default="engagement-letter")
    args = parser.parse_args()
    output = Path(args.output).expanduser().resolve()
    if output.suffix.lower() != ".docx":
        parser.error("output must be a .docx file")
    build(output, args.profile)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
