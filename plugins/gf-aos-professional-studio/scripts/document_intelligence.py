#!/usr/bin/env python3
"""Extract document content into a separate GF-AOS workspace without changing sources."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

SUPPORTED = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "spreadsheet",
    ".xlsm": "spreadsheet",
    ".pptx": "presentation",
    ".txt": "text",
    ".md": "text",
    ".csv": "csv",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tif": "image",
    ".tiff": "image",
}

MAX_OOXML_UNCOMPRESSED = 500 * 1024 * 1024
MAX_OOXML_RATIO = 200


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_ooxml_container(path: Path) -> None:
    """Reject suspiciously large or compressed Office containers before parsing."""
    try:
        with zipfile.ZipFile(path) as archive:
            total = 0
            for item in archive.infolist():
                total += item.file_size
                if total > MAX_OOXML_UNCOMPRESSED:
                    raise ValueError("Archivio Office oltre il limite non compresso")
                if item.file_size and item.compress_size == 0:
                    raise ValueError("Voce Office con compressione non valida")
                if item.compress_size and item.file_size / item.compress_size > MAX_OOXML_RATIO:
                    raise ValueError("Rapporto di compressione Office sospetto")
                if item.flag_bits & 0x1:
                    raise ValueError("Archivio Office cifrato non supportato")
    except zipfile.BadZipFile as exc:
        raise ValueError("Archivio Office non valido") from exc


def extract_pdf(path: Path) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("PDF cifrato non estraibile") from exc
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"\n## Pagina {index}\n\n{page.extract_text() or ''}")
    metadata = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
    return "".join(pages).strip(), {"pages": len(reader.pages), "pdf_metadata": metadata}


def extract_docx(path: Path) -> tuple[str, dict]:
    from docx import Document

    validate_ooxml_container(path)
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    table_cells = 0
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n## Tabella {table_index}\n")
        for row in table.rows:
            values = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            table_cells += len(values)
            parts.append(" | ".join(values))
    props = document.core_properties
    metadata = {
        "title": props.title or "",
        "author": props.author or "",
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
    }
    return "\n".join(parts).strip(), {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "table_cells": table_cells,
        "docx_metadata": metadata,
    }


def extract_spreadsheet(path: Path) -> tuple[str, dict]:
    from openpyxl import load_workbook

    validate_ooxml_container(path)
    workbook = load_workbook(
        filename=str(path), read_only=True, data_only=False, keep_links=False
    )
    parts, sheets, non_empty, formulas = [], [], 0, 0
    for sheet in workbook.worksheets:
        sheets.append(sheet.title)
        parts.append(f"\n## Foglio: {sheet.title}\n\n```tsv")
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                if cell.value is None:
                    values.append("")
                    continue
                value = str(cell.value)
                non_empty += 1
                formulas += int(value.startswith("="))
                values.append(value.replace("\n", " "))
            if any(values):
                parts.append("\t".join(values))
        parts.append("```")
    workbook.close()
    return "\n".join(parts).strip(), {
        "sheets": sheets,
        "non_empty_cells": non_empty,
        "formula_cells": formulas,
        "formula_note": "Le formule sono registrate ma non calcolate dal motore.",
    }


def extract_presentation(path: Path) -> tuple[str, dict]:
    from pptx import Presentation

    validate_ooxml_container(path)
    presentation = Presentation(str(path))
    parts, text_shapes, tables = [], 0, 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"\n## Diapositiva {slide_index}\n")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_shapes += 1
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                tables += 1
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip(), {
        "slides": len(presentation.slides),
        "text_shapes": text_shapes,
        "tables": tables,
    }


def extract_text(path: Path) -> tuple[str, dict]:
    return path.read_text(encoding="utf-8", errors="replace"), {}


def extract_csv(path: Path) -> tuple[str, dict]:
    lines, rows = ["```tsv"], 0
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for row in csv.reader(handle, dialect):
            rows += 1
            lines.append("\t".join(value.replace("\n", " ") for value in row))
    lines.append("```")
    return "\n".join(lines), {"rows": rows, "delimiter": dialect.delimiter}


def extract_ocr(path: Path, language: str) -> tuple[str, dict]:
    executable = shutil.which("tesseract")
    if not executable:
        raise RuntimeError("Tesseract non disponibile; OCR non eseguito")
    languages_result = subprocess.run(
        [executable, "--list-langs"], check=False, capture_output=True, text=True, timeout=30
    )
    available = {
        item.strip()
        for item in languages_result.stdout.splitlines()
        if item.strip() and not item.lower().startswith("list of available")
    }
    requested = [item for item in language.split("+") if item]
    selected = [item for item in requested if item in available]
    missing = [item for item in requested if item not in available]
    if not selected:
        raise RuntimeError(f"Lingue OCR non disponibili: {', '.join(requested)}")
    used_language = "+".join(selected)
    result = subprocess.run(
        [executable, str(path), "stdout", "-l", used_language],
        check=False,
        capture_output=True,
        text=True,
        timeout=300,
    )
    if result.returncode:
        message = result.stderr.strip().splitlines()[-1] if result.stderr.strip() else "errore OCR"
        raise RuntimeError(message)
    return result.stdout, {
        "ocr_language_requested": language,
        "ocr_language_used": used_language,
        "ocr_languages_missing": missing,
        "ocr_engine": "tesseract",
    }


EXTRACTORS: dict[str, Callable[[Path], tuple[str, dict]]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_spreadsheet,
    ".xlsm": extract_spreadsheet,
    ".pptx": extract_presentation,
    ".txt": extract_text,
    ".md": extract_text,
    ".csv": extract_csv,
}


def source_files(source: Path) -> list[Path]:
    if source.is_file():
        return [source]
    files = []
    for root, directories, names in os.walk(source, followlinks=False):
        directories[:] = sorted(
            name for name in directories if not (Path(root) / name).is_symlink()
        )
        for name in sorted(names):
            candidate = Path(root) / name
            if not candidate.is_symlink() and candidate.suffix.lower() in SUPPORTED:
                files.append(candidate)
    return files


def safe_locator(relative: str, digest: str) -> str:
    stem = Path(relative).stem[:60]
    cleaned = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in stem)
    return f"markdown/{cleaned or 'documento'}-{digest[:12]}.md"


def atomic_text_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", dir=path.parent, delete=False
    ) as handle:
        handle.write(content)
        temporary = Path(handle.name)
    temporary.replace(path)


def build_index(
    source: Path,
    output_dir: Path,
    *,
    ocr: bool = False,
    ocr_language: str = "ita+eng",
    max_chars: int = 1_000_000,
    max_file_bytes: int = 200 * 1024 * 1024,
) -> dict:
    if source.is_symlink():
        raise ValueError("Sorgente simbolica non consentita")
    if output_dir.is_symlink():
        raise ValueError("Workspace simbolico non consentito")
    source, output_dir = source.resolve(), output_dir.resolve()
    if not source.exists():
        raise ValueError("La sorgente non esiste")
    if source.is_file() and source.suffix.lower() not in SUPPORTED:
        raise ValueError("Formato sorgente non supportato")
    if source.is_dir() and (output_dir == source or source in output_dir.parents):
        raise ValueError("Il workspace di output deve essere esterno alla cartella sorgente")
    output_dir.mkdir(parents=True, exist_ok=True)
    text_dir = output_dir / "markdown"
    if text_dir.exists() and text_dir.is_symlink():
        raise ValueError("Cartella Markdown simbolica non consentita")

    base = source if source.is_dir() else source.parent
    records = []
    for path in source_files(source):
        relative = str(path.relative_to(base))
        stat = path.stat()
        digest = sha256_file(path)
        record = {
            "path": relative,
            "type": SUPPORTED.get(path.suffix.lower(), "unsupported"),
            "bytes": stat.st_size,
            "modified_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "sha256": digest,
            "status": "estratto",
            "extractor": None,
            "text_chars": 0,
            "text_locator": None,
            "text_format": "markdown",
            "truncated": False,
            "details": {},
            "issues": [],
        }
        try:
            if stat.st_size > max_file_bytes:
                raise ValueError(
                    f"File oltre il limite di {max_file_bytes // (1024 * 1024)} MiB"
                )
            suffix = path.suffix.lower()
            if suffix in EXTRACTORS:
                extractor = EXTRACTORS[suffix]
                text, details = extractor(path)
                record["extractor"] = extractor.__name__
            elif SUPPORTED.get(suffix) == "image" and ocr:
                text, details = extract_ocr(path, ocr_language)
                record["extractor"] = "tesseract"
            else:
                text, details = "", {}
                record["status"] = "ocr_non_richiesto"
                record["issues"].append("Immagine rilevata; usare --ocr per estrarre il testo")
            if len(text) > max_chars:
                text = text[:max_chars]
                record["truncated"] = True
                record["issues"].append(f"Testo limitato a {max_chars} caratteri")
            if text:
                locator = safe_locator(relative, digest)
                atomic_text_write(output_dir / locator, text)
                record["text_locator"] = locator
                record["text_chars"] = len(text)
            record["details"] = details
            missing_languages = details.get("ocr_languages_missing", [])
            if missing_languages:
                record["issues"].append(
                    "Lingue OCR non disponibili: " + ", ".join(missing_languages)
                )
        except Exception as exc:
            record["status"] = "errore"
            record["issues"].append(f"{type(exc).__name__}: {exc}")
        records.append(record)

    result = {
        "schema_version": "1.0",
        "created_utc": utc_now(),
        "source": str(source),
        "source_mode": "read_only",
        "output_dir": str(output_dir),
        "ocr_requested": ocr,
        "count": len(records),
        "extracted": sum(item["status"] == "estratto" for item in records),
        "errors": sum(item["status"] == "errore" for item in records),
        "records": records,
        "limitations": [
            "L'estrazione non prova autenticità, completezza o approvazione del documento.",
            "Formule, macro, firme digitali e collegamenti esterni non sono eseguiti o validati.",
            "OCR e impaginazione complessa richiedono controllo umano sul documento originale.",
        ],
    }
    atomic_text_write(
        output_dir / "document_index.json",
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
    )
    rows = [
        "# Registro Document Intelligence GF-AOS",
        "",
        f"Documenti rilevati: {result['count']}. Estratti: {result['extracted']}. Errori: {result['errors']}.",
        "",
        "| Documento | Tipo | Stato | Caratteri | Hash SHA-256 | Evidenza |",
        "| --- | --- | --- | ---: | --- | --- |",
    ]
    for item in records:
        display_path = item["path"].replace("|", "¦")
        locator = item["text_locator"] or "—"
        rows.append(
            f"| {display_path} | {item['type']} | {item['status']} | "
            f"{item['text_chars']} | `{item['sha256'][:16]}…` | {locator} |"
        )
    rows.extend(["", "## Limiti", ""] + [f"- {item}" for item in result["limitations"]])
    atomic_text_write(output_dir / "document_index.md", "\n".join(rows) + "\n")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", help="File o cartella sorgente")
    parser.add_argument("--output-dir", required=True, help="Workspace separato di output")
    parser.add_argument("--ocr", action="store_true", help="Abilita OCR immagini tramite Tesseract")
    parser.add_argument("--ocr-language", default="ita+eng")
    parser.add_argument("--max-chars", type=int, default=1_000_000)
    parser.add_argument("--max-file-mb", type=int, default=200)
    args = parser.parse_args()
    if args.max_chars < 1:
        parser.error("--max-chars deve essere positivo")
    if args.max_file_mb < 1:
        parser.error("--max-file-mb deve essere positivo")
    result = build_index(
        Path(args.source),
        Path(args.output_dir),
        ocr=args.ocr,
        ocr_language=args.ocr_language,
        max_chars=args.max_chars,
        max_file_bytes=args.max_file_mb * 1024 * 1024,
    )
    print(
        f"Analizzati {result['count']} documenti: "
        f"{result['extracted']} estratti, {result['errors']} errori"
    )
    return 0 if result["errors"] == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
